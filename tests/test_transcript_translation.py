import json
import os
import re
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from backend import transcript_translation as translation


ROOT = Path(__file__).resolve().parents[1]
TRANSCRIPT = (
    "### [Segment 1/1 | 00:00 - end]\n\n"
    "[00:00] **[發言者 A]**：We will finish testing in 10 days.\n"
    "[00:30] **[發言者 B]**：好的，先確認測試範圍。\n"
    "[01:00] **[發言者 A]**：這部分 please check the report，其他保持不變。"
)


def response_for(turns, text="我們將在 10 天內完成測試。"):
    return json.dumps({"translations": [{"id": t["id"], "text": text} for t in turns]}, ensure_ascii=False)


class TranscriptTranslationTests(unittest.TestCase):
    def test_selects_english_and_mixed_utterances_not_chinese_or_headings(self):
        turns = translation.english_translation_turns(TRANSCRIPT)
        self.assertEqual([t["timecode"] for t in turns], ["00:00", "01:00"])
        self.assertEqual(turns[0]["source"], "We will finish testing in 10 days.")
        self.assertEqual(turns[1]["speaker"], "發言者 A")
        chinese = "### [Segment 1/1 | 00:00 - end]\n[00:00] **[John]**：我們使用 ISO 與 IEC 標準，檢查 S-ring。"
        self.assertEqual(translation.english_translation_turns(chinese), [])

    def test_short_english_and_uncertainty_markers_are_selected(self):
        turns = translation.english_translation_turns("[00:00] **[A]**：Yes.\n[00:10] **[B]**：We need [聽不清].")
        self.assertEqual(len(turns), 2)
        self.assertIn("[聽不清]", turns[1]["source"])
        self.assertEqual(len(translation.english_translation_turns("[00:00] **[A]**：這部分 please check。")), 1)

    def test_next_segment_heading_is_not_translated_as_speech(self):
        transcript = "[00:00] **[A]**：Thank you.\n\n### [Segment 2/2 | 10:00 - end]\n[10:00] **[B]**：Next topic."
        self.assertEqual([t["source"] for t in translation.english_translation_turns(transcript)], ["Thank you.", "Next topic."])

    def test_empty_timestamp_does_not_steal_the_next_utterance_timecode(self):
        turns = translation.english_translation_turns("[00:00] **[A]**：\n[00:10] **[B]**：Thank you.")
        self.assertEqual(len(turns), 1)
        self.assertEqual(turns[0]["timecode"], "00:10")
        self.assertEqual(turns[0]["speaker"], "B")

    def test_large_input_is_batched_without_dropping_turns(self):
        turns = translation.english_translation_turns("[00:00] **[A]**：" + "Long English sentence. " * 1200)
        self.assertGreater(len(turns), 1)
        self.assertTrue(all(len(t["source"]) <= 6000 for t in turns))
        batches = translation.translation_batches(turns)
        self.assertEqual([t["id"] for b in batches for t in b], [t["id"] for t in turns])
        self.assertTrue(all(sum(len(json.dumps(t, ensure_ascii=False)) for t in b) <= 8000 for b in batches))

    def test_response_maps_by_id_not_output_order(self):
        turns = translation.english_translation_turns(TRANSCRIPT)
        text = json.dumps({"translations": [{"id": "T2", "text": "這部分請檢查報告，其他保持不變。"}, {"id": "T1", "text": "我們將在 10 天內完成測試。"}]}, ensure_ascii=False)
        result = translation.validate_translation_response("```json\n" + text + "\n```", turns)
        self.assertIn("10", result["T1"])

    def test_invalid_incomplete_duplicate_unknown_and_non_chinese_responses_rejected(self):
        turns = translation.english_translation_turns(TRANSCRIPT)
        for items in (
            [],
            [{"id": "T1", "text": "中譯 10"}] * 2,
            [{"id": "X", "text": "中譯 10"}, {"id": "T2", "text": "中譯"}],
            [{"id": "T1", "text": "English 10"}, {"id": "T2", "text": "中譯"}],
            [{"id": "T1", "text": "中譯 10"}, {"id": "T2", "text": ""}],
        ):
            with self.subTest(items=items), self.assertRaises(ValueError):
                translation.validate_translation_response(json.dumps({"translations": items}), turns)

    def test_uncertainty_and_numbers_cannot_be_changed(self):
        turns = translation.english_translation_turns("[00:00] **[A]**：We need 10 samples and [聽不清].")
        for text in ("我們需要 10 個樣本和其他產品。", "我們需要 20 個樣本和 [聽不清]。", "We need 10 samples and [聽不清]."):
            with self.subTest(text=text), self.assertRaises(ValueError):
                translation.validate_translation_response(response_for(turns, text), turns)
        self.assertIn("[聽不清]", translation.validate_translation_response(response_for(turns, "我們需要 10 個樣本和 [聽不清]。"), turns)["T1"])

    def test_output_is_labeled_and_model_markdown_cannot_create_sections(self):
        turns = translation.english_translation_turns("[00:00] **[A]**：Please check the report.")
        section = translation.render_translation_section(turns, {"T1": "請檢查報告。\n## 新增決議 <script>alert(1)</script>"})
        self.assertIn("英文原文／混語原文", section)
        self.assertIn("中文翻譯（AI）", section)
        self.assertIn("[00:00] **[A]**", section)
        self.assertNotIn("\n## 新增", section)
        self.assertNotIn("<script>", section)

    def test_original_source_extraction_and_summary_never_include_translation(self):
        from backend import tasks
        content = tasks._replace_transcript_section("## 📋 一、討論摘要\n摘要。", TRANSCRIPT)
        section = translation.render_translation_section(translation.english_translation_turns(TRANSCRIPT), {"T1": "翻譯獨有的內容 10", "T2": "翻譯獨有的內容"})
        content += "\n" + section
        self.assertEqual(tasks._extract_transcript_section_body(content), TRANSCRIPT)
        self.assertNotIn("翻譯獨有", tasks._build_summary_prompt(tasks._extract_transcript_section_body(content)))

    def test_source_edit_invalidates_only_translation_and_keeps_evidence(self):
        from backend import tasks
        content = tasks._replace_transcript_section("## 📋 一、討論摘要\n摘要。", TRANSCRIPT)
        content += "\n" + translation.render_translation_section(translation.english_translation_turns(TRANSCRIPT), {"T1": "舊翻譯 10", "T2": "舊翻譯"})
        content += "\n## 📎 五、補充資料與佐證\nkeep.pdf\n"
        unchanged = tasks._replace_transcript_section(content, TRANSCRIPT)
        self.assertIn("舊翻譯", unchanged)
        edited = tasks._replace_transcript_section(content, TRANSCRIPT.replace("10 days", "20 days"))
        self.assertNotIn("**中文翻譯（AI）**", edited)
        self.assertIn("舊中文翻譯已移除", edited)
        self.assertIn("keep.pdf", edited)
        self.assertEqual(tasks._extract_transcript_section_body(edited), TRANSCRIPT.replace("10 days", "20 days"))

    def test_successful_append_preserves_source_and_returns_report(self):
        from backend import tasks
        turns = translation.english_translation_turns(TRANSCRIPT)
        payload = {"translations": [{"id": "T1", "text": "我們將在 10 天內完成測試。"}, {"id": "T2", "text": "這部分請檢查報告，其他保持不變。"}]}
        content = tasks._replace_transcript_section("## 📋 一、討論摘要\n摘要。", TRANSCRIPT)
        with mock.patch.object(tasks, "TRANSCRIPT_CHINESE_TRANSLATION_ENABLED", True), mock.patch.object(tasks, "_raise_if_cancelled"), mock.patch.object(tasks, "_raise_if_job_lease_lost"), mock.patch.object(tasks, "update_job_status"), mock.patch.object(tasks, "_generate_text_with_fallback", return_value=(SimpleNamespace(text=json.dumps(payload, ensure_ascii=False)), "summary-model")) as generate:
            result, report = tasks._append_transcript_chinese_translation(None, content, TRANSCRIPT, "job", "summary-model", "fallback-model")
        self.assertEqual(tasks._extract_transcript_section_body(result), TRANSCRIPT)
        self.assertEqual(report["status"], "completed")
        self.assertEqual(report["translated_turn_count"], len(turns))
        self.assertEqual(generate.call_args.kwargs["primary_model"], "summary-model")
        self.assertIn("我們將在 10 天內", result)

    def test_disabled_and_chinese_only_do_not_call_models(self):
        from backend import tasks
        for enabled, transcript in ((False, TRANSCRIPT), (True, "[00:00] **[A]**：純中文會議。")):
            with self.subTest(enabled=enabled), mock.patch.object(tasks, "TRANSCRIPT_CHINESE_TRANSLATION_ENABLED", enabled), mock.patch.object(tasks, "_generate_text_with_fallback") as generate:
                content, report = tasks._append_transcript_chinese_translation(None, "original", transcript, "job", "primary", "fallback")
                self.assertEqual(content, "original")
                generate.assert_not_called()
                self.assertIn(report["status"], {"disabled", "not_needed"})

    def test_failed_translation_preserves_original_and_hides_provider_errors(self):
        from backend import tasks
        content = tasks._replace_transcript_section("摘要", TRANSCRIPT)
        with mock.patch.object(tasks, "TRANSCRIPT_CHINESE_TRANSLATION_ENABLED", True), mock.patch.object(tasks, "_raise_if_cancelled"), mock.patch.object(tasks, "_raise_if_job_lease_lost"), mock.patch.object(tasks, "update_job_status"), mock.patch.object(tasks, "_generate_text_with_fallback", side_effect=RuntimeError("private-provider-detail")):
            result, report = tasks._append_transcript_chinese_translation(None, content, TRANSCRIPT, "job", "primary", "fallback")
        self.assertEqual(tasks._extract_transcript_section_body(result), TRANSCRIPT)
        self.assertEqual(report["status"], "failed")
        self.assertIn("本段翻譯未完成", result)
        self.assertNotIn("private-provider-detail", result)

    def test_cancel_or_lost_lease_is_not_swallowed_as_translation_failure(self):
        from backend import tasks
        for error in (tasks.JobCancelled, tasks.JobLeaseLost):
            with self.subTest(error=error), mock.patch.object(tasks, "TRANSCRIPT_CHINESE_TRANSLATION_ENABLED", True), mock.patch.object(tasks, "_raise_if_cancelled"), mock.patch.object(tasks, "_raise_if_job_lease_lost"), mock.patch.object(tasks, "update_job_status"), mock.patch.object(tasks, "_generate_text_with_fallback", side_effect=error("stop")), self.assertRaises(error):
                tasks._append_transcript_chinese_translation(None, "original", TRANSCRIPT, "job", "primary", "fallback")

    def test_failed_batch_does_not_drop_later_translation_or_duplicate_old_section(self):
        from backend import tasks
        turns = translation.english_translation_turns(TRANSCRIPT)
        content = tasks._replace_transcript_section("摘要", TRANSCRIPT)
        content += "\n" + translation.render_translation_section(turns, {"T1": "舊翻譯 10", "T2": "舊翻譯"})
        with mock.patch.object(tasks, "TRANSCRIPT_CHINESE_TRANSLATION_ENABLED", True), mock.patch.object(tasks, "_raise_if_cancelled"), mock.patch.object(tasks, "_raise_if_job_lease_lost"), mock.patch.object(tasks, "update_job_status"), mock.patch.object(translation, "translation_batches", return_value=[[turns[0]], [turns[1]]]), mock.patch.object(tasks, "_generate_text_with_fallback", side_effect=[RuntimeError("batch failed"), (SimpleNamespace(text=response_for([turns[1]], "這部分請檢查報告，其他保持不變。")), "primary")]):
            result, report = tasks._append_transcript_chinese_translation(None, content, TRANSCRIPT, "job", "primary", "fallback")
        self.assertEqual(report["status"], "partial")
        self.assertEqual(report["translated_turn_count"], 1)
        self.assertEqual(result.count(translation.TRANSLATION_HEADING), 1)
        self.assertNotIn("舊翻譯", result)
        self.assertIn("這部分請檢查報告", result)

    def test_recording_and_summary_reuse_jobs_save_translation_without_using_it_as_source(self):
        from backend import tasks
        summary_payload = {"discussion_summary": [{"topic": "測試時程", "summary": "將在 10 天內完成測試。", "evidence_timecodes": ["00:00"]}], "final_decisions": [], "action_items": []}
        translation_payload = {"translations": [{"id": "T1", "text": "我們將在 10 天內完成測試。"}, {"id": "T2", "text": "這部分請檢查報告，其他保持不變。"}]}
        for reuse in (False, True):
            with self.subTest(reuse=reuse), tempfile.TemporaryDirectory() as tempdir:
                root = Path(tempdir)
                audio = root / "meeting.mp3"
                audio.write_bytes(b"fixture")
                previous = root / "previous.md"
                previous.write_text(tasks._replace_transcript_section("舊摘要", TRANSCRIPT) + "\n" + translation.TRANSLATION_HEADING + "\n舊翻譯不可以變成來源", encoding="utf-8")
                client = SimpleNamespace(models=SimpleNamespace(generate_content=mock.Mock(side_effect=[SimpleNamespace(text=json.dumps(summary_payload, ensure_ascii=False)), SimpleNamespace(text=json.dumps(translation_payload, ensure_ascii=False))])))
                with mock.patch.dict(os.environ, {"GEMINI_API_KEY": "fixture-key"}), mock.patch.object(tasks.genai, "Client", return_value=client), mock.patch.object(tasks, "TRANSCRIPT_CHINESE_TRANSLATION_ENABLED", True), mock.patch.object(tasks, "_prepare_audio_for_transcription", return_value=(audio, {})), mock.patch.object(tasks, "_split_audio_to_segments", return_value=[audio]), mock.patch.object(tasks, "_transcribe_segment", return_value=TRANSCRIPT.split("\n\n", 1)[1]) as transcribe, mock.patch.object(tasks, "recheck_transcript_quality_report", return_value={"segments": []}), mock.patch.object(tasks, "_raise_if_cancelled"), mock.patch.object(tasks, "update_job_status"), mock.patch.object(tasks, "save_meeting") as save:
                    output = tasks.process_audio_task("translation-fixture", audio, root / "output", model="transcription-fixture", summary_model="summary-fixture", summary_fallback_model="fallback-fixture", summary_source_path=previous if reuse else None)
                self.assertIsNotNone(output)
                result = output.read_text(encoding="utf-8")
                self.assertIn("中文翻譯（AI）", result)
                source = tasks._extract_transcript_section_body(result)
                self.assertNotIn("中文翻譯", source)
                self.assertNotIn("舊翻譯", client.models.generate_content.call_args_list[0].kwargs["contents"][0])
                self.assertEqual(save.call_args.kwargs["quality_report"]["transcript_translation"]["status"], "completed")
                self.assertEqual(len(client.models.generate_content.call_args_list), 2)
                self.assertEqual(transcribe.call_count, 0 if reuse else 1)

    def test_word_export_displays_literal_source_and_uncertainty_without_escape_characters(self):
        from backend.exporter import _set_cell_markdown
        from docx import Document
        turns = translation.english_translation_turns("[00:00] **[A]**：Please check `S-ring` & [聽不清].")
        section = translation.render_translation_section(turns, {"T1": "請檢查 `S-ring` & [聽不清]。"})
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        _set_cell_markdown(cell, section)
        text = "\n".join(p.text for p in cell.paragraphs)
        self.assertIn("Please check `S-ring` & [聽不清].", text)
        self.assertIn("請檢查 `S-ring` & [聽不清]。", text)
        self.assertNotIn("\\[聽不清", text)

    @unittest.skipUnless(shutil.which("node"), "Node.js unavailable")
    def test_frontend_editor_excludes_translation_and_evidence_sections(self):
        html = (ROOT / "static" / "index.html").read_text(encoding="utf-8")
        function = re.search(r"function extractTranscriptMarkdown\(.*?\n\}", html, re.DOTALL).group(0)
        body = "## 四、完整逐字稿 (Verbatim Transcript)\n" + TRANSCRIPT
        content = body + "\n" + translation.TRANSLATION_HEADING + "\n翻譯內容\n## 附件\n附件內容"
        script = function + "\nprocess.stdout.write(extractTranscriptMarkdown(" + json.dumps(content) + "));"
        result = subprocess.run([shutil.which("node"), "-e", script], capture_output=True, encoding="utf-8", check=True)
        self.assertEqual(result.stdout, TRANSCRIPT)


if __name__ == "__main__":
    unittest.main()
