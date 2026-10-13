"""Topic reader/export contracts; synthetic fixtures only, no live writes or AI."""
from copy import deepcopy
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest import mock

from backend.minutes_view import build_minutes_view, topic_minutes_markdown
from backend.structured_minutes import parse_standardized_minutes
from tests.test_structured_backfill import MINUTES


CONTENT = MINUTES + '[00:30] **[A]**：Keep the English original.\n\n## 英文中譯對照（AI 翻譯，僅供參考）\n保留英文原文。\n'


def record():
    return {"id": 1, "title": "Reader fixture", "date": "2026/10/08", "full_content": CONTENT,
            "review_status": "generated", "source_audio": "test.wav", "output_path": "not-a-live-file",
            "created_at": "2026-10-08T00:00:00", "structured_summary": parse_standardized_minutes(CONTENT),
            "structured_items": [{"item_key": "R1", "review_status": "reviewed"}],
            "quality_report": {}, "source_media_available": False}


class MinutesViewTests(unittest.TestCase):
    def test_read_model_does_not_mutate_source(self):
        r = record(); before = deepcopy(r)
        view = build_minutes_view(r)
        self.assertTrue(view["available"])
        self.assertEqual(r, before)
        self.assertEqual(view["decisions"][0]["review_status"], "reviewed")
        self.assertEqual(view["review_status"], "generated")

    def test_missing_snapshot_uses_deterministic_parser(self):
        r = record(); r.pop("structured_summary")
        self.assertEqual(build_minutes_view(r)["source"], "markdown")

    def test_manual_confirmation_is_not_replaced_by_stale_markdown(self):
        r = record(); r["structured_summary"]["action_items"][0]["owner"] = "具名負責人"
        self.assertEqual(build_minutes_view(r)["actions"][0]["owner"], "具名負責人")
        self.assertIn("具名負責人", topic_minutes_markdown(r))

    def test_original_and_translation_remain_intact_in_appendix(self):
        r = record(); v = build_minutes_view(r)
        self.assertTrue(CONTENT.endswith(v["appendix_markdown"]))
        self.assertTrue(topic_minutes_markdown(r).endswith(v["appendix_markdown"]))

    def test_unassigned_and_shared_relations_are_not_guessed(self):
        r = record(); s = r["structured_summary"]
        d2 = deepcopy(s["discussion_summary"][0]); d2["id"] = "D2"; s["discussion_summary"].append(d2)
        s["final_decisions"][0]["related_discussions"] = ["D1", "D2"]
        s["action_items"][0]["related_discussions"] = ["D404"]
        v = build_minutes_view(r)
        self.assertEqual(v["unassigned"], ["A1"])
        self.assertEqual(v["actions"][0]["related_discussions"], ["D404"])
        self.assertEqual(v["actions"][0]["related_decisions"], ["R1"])
        self.assertIn("未明確歸屬議題", topic_minutes_markdown(r))

    def test_identical_actions_are_not_merged(self):
        r = record(); a = deepcopy(r["structured_summary"]["action_items"][0]); a["id"] = "A2"
        r["structured_summary"]["action_items"].append(a)
        self.assertEqual(len(build_minutes_view(r)["actions"]), 2)

    def test_pending_decision_condition_is_retained(self):
        r = record(); d = r["structured_summary"]["final_decisions"][0]
        d.update(status="pending", basis="只有測試通過才發包")
        output = topic_minutes_markdown(r)
        self.assertIn("待確認", output); self.assertIn("只有測試通過才發包", output)

    def test_review_status_is_never_implicitly_approved(self):
        r = record(); v = build_minutes_view(r)
        self.assertEqual(v["review_status"], "generated")
        self.assertIn("尚未複核", topic_minutes_markdown(r))
        r["review_status"] = "approved"
        self.assertIn("文件審查狀態：已核准", topic_minutes_markdown(r))

    def test_source_warning_follows_actual_time_range_not_fixture_id(self):
        r = record(); r["quality_review_segment_details"] = [{"start_seconds": 15, "end_seconds": 25, "label": "需複核"}]
        v = build_minutes_view(r)
        self.assertTrue(v["decisions"][0]["review_ranges"])
        self.assertFalse(v["discussions"][0]["review_ranges"])
        self.assertFalse(v["actions"][0]["review_ranges"])

    def test_hours_timecode_supported(self):
        r = record(); r["structured_summary"]["action_items"][0]["source_timecodes"] = ["01:02:03"]
        r["quality_review_segment_details"] = [{"start_seconds": 3720, "end_seconds": 3730}]
        self.assertTrue(build_minutes_view(r)["actions"][0]["review_ranges"])

    def test_unsupported_legacy_keeps_original(self):
        r = record(); r["full_content"] = "# 舊紀錄\n重要但沒有 D/R/A 的原始內容"; r.pop("structured_summary")
        self.assertFalse(build_minutes_view(r)["available"])
        self.assertEqual(topic_minutes_markdown(r), r["full_content"])

    def test_malformed_snapshot_does_not_silently_fall_back(self):
        for bad in [{}, [], "bad", {"discussion_summary": None}]:
            r = record(); r["structured_summary"] = bad
            self.assertFalse(build_minutes_view(r)["available"])

    def test_snapshot_with_parse_errors_falls_back_to_original(self):
        r = record(); r["structured_summary"]["parse_errors"] = ["invalid action row"]
        self.assertFalse(build_minutes_view(r)["available"])
        self.assertEqual(topic_minutes_markdown(r), r["full_content"])

    def test_export_keeps_item_review_and_source_warnings(self):
        r = record(); r["quality_review_segment_details"] = [{"start_seconds": 0, "end_seconds": 100}]
        r["structured_summary"]["action_items"][0].update(owner="發言者 A", due="未提及")
        output = topic_minutes_markdown(r)
        for warning in ("人工審查：已人工複核", "來源需複核", "負責人姓名待確認", "期限待確認"):
            self.assertIn(warning, output)

    def test_duplicate_or_unsafe_ids_fall_back(self):
        for key in ["D1", '<img onerror="bad">', "R1' onclick='bad"]:
            r = record(); r["structured_summary"]["final_decisions"][0]["id"] = key
            self.assertFalse(build_minutes_view(r)["available"])

    def test_empty_topic_list_falls_back(self):
        r = record(); r["structured_summary"]["discussion_summary"] = []
        self.assertFalse(build_minutes_view(r)["available"])

    def test_empty_decisions_and_actions_are_supported(self):
        r = record(); r["structured_summary"].update(final_decisions=[], action_items=[])
        self.assertTrue(build_minutes_view(r)["available"])
        self.assertIn("原稿未列明", topic_minutes_markdown(r))

    def test_extra_middle_section_is_not_dropped(self):
        r = record(); r["full_content"] = CONTENT.replace("## 二、", "## 自訂補充\n重要的補充\n\n## 二、")
        self.assertFalse(build_minutes_view(r)["available"])
        self.assertEqual(topic_minutes_markdown(r), r["full_content"])

    def test_handwritten_prose_or_unknown_fields_are_not_dropped(self):
        for extra in ["會後補充這個條件不可刪除", "- 特別承諾：必須保留"]:
            r = record(); r["full_content"] = CONTENT.replace("- 摘要：", extra + "\n- 摘要：")
            self.assertFalse(build_minutes_view(r)["available"])
            self.assertEqual(topic_minutes_markdown(r), r["full_content"])

    def test_unparsed_custom_table_row_is_not_dropped(self):
        r = record(); r["full_content"] = CONTENT.replace("| R1 |", "| 特殊補充 | D1 | 不可遺失 | 有條件 | pending |\n| R1 |")
        self.assertFalse(build_minutes_view(r)["available"])
        self.assertEqual(topic_minutes_markdown(r), r["full_content"])

    def test_frontmatter_previous_minutes_and_extra_appendix_preserved(self):
        r = record(); r["full_content"] = "---\ntitle: 測試\n---\n## 前次紀錄追蹤\n前次尚未結案\n" + CONTENT + "\n## 補充佐證\n附件 A\n"
        output = topic_minutes_markdown(r)
        self.assertIn("前次尚未結案", output); self.assertTrue(output.endswith("附件 A\n"))

    def test_export_table_keeps_columns_with_pipes_and_newlines(self):
        r = record(); r["structured_summary"]["action_items"][0]["task"] = "A|B\n下一行"
        self.assertIn("A｜B<br>下一行", topic_minutes_markdown(r))

    def test_node_renderer_contracts(self):
        node = shutil.which("node")
        if not node: self.skipTest("Node.js is unavailable")
        result = subprocess.run([node, "tests/minutes_reader_test.js"], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


class MinutesViewIntegrationTests(unittest.TestCase):
    def test_export_route_default_and_explicit_layouts(self):
        from backend import main
        from tests.http_helpers import asgi_request
        r = record()
        with tempfile.TemporaryDirectory() as directory:
            def export(data, destination, *, layout):
                Path(destination).write_bytes(b"test export route response")
                return True
            with mock.patch.object(main, "get_meeting", return_value=r), \
                 mock.patch.object(main, "Path", return_value=Path(directory)), \
                 mock.patch.object(main, "apply_quality_preview_fields", return_value=r) as quality, \
                 mock.patch.object(main, "export_meeting_to_docx", side_effect=export) as generate:
                for query, expected in (("", "topics"), ("?layout=topics", "topics"), ("?layout=original", "original")):
                    response = asgi_request(main.app, "GET", "/meetings/1/export/docx" + query)
                    self.assertEqual(response.status_code, 200, response.text)
                    self.assertEqual(generate.call_args.kwargs["layout"], expected)
                    self.assertIn(".docx", response.headers["content-disposition"])
                self.assertEqual(quality.call_count, 2)

    def test_detail_exposes_projection_without_database_updates(self):
        from backend import main
        from tests.http_helpers import asgi_request
        r = record()
        with mock.patch.object(main, "get_meeting", return_value=r), mock.patch.object(main, "update_meeting_quality_report") as update:
            response = asgi_request(main.app, "GET", "/meetings/1")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()["minutes_view"]["available"])
        self.assertEqual(response.json()["minutes_view"]["actions"][0]["id"], "A1")
        update.assert_not_called()

    def test_word_export_both_layouts_preserve_template_and_transcript(self):
        from docx import Document
        from backend.exporter import export_meeting_to_docx
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "template.docx"
            doc = Document(); table = doc.add_table(rows=6, cols=2)
            table.cell(0,0).text = "受控範本表頭"; doc.save(template)
            for layout in ("topics", "original"):
                output = Path(directory) / (layout + ".docx")
                with mock.patch.dict(os.environ, {"MEETING_DOCX_TEMPLATE_PATH":str(template)}):
                    self.assertTrue(export_meeting_to_docx(record(), str(output), layout=layout))
                result = Document(output)
                text = "\n".join(p.text for p in result.tables[0].cell(5,0).paragraphs)
                self.assertEqual(result.tables[0].cell(0,0).text, "受控範本表頭")
                self.assertIn("Keep the English original.", text); self.assertIn("保留英文原文。", text)
                self.assertEqual("議題式會議紀錄" in text, layout == "topics")

    def test_invalid_export_layout_rejected(self):
        from backend import main
        from tests.http_helpers import asgi_request
        with mock.patch.object(main, "get_meeting") as get:
            response = asgi_request(main.app, "GET", "/meetings/1/export/docx?layout=unsafe")
        self.assertEqual(response.status_code, 422)
        get.assert_not_called()


if __name__ == "__main__":
    unittest.main()
